import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

def build_presentation(input_pptx="SIH2026-IDEA-Presentation-Format.pptx",
                       output_pptx="SIH2026_KSI_Submission.pptx",
                       arch_img="workflow_architecture.png"):
    """
    Overhauls and populates the presentation deck adhering strictly to:
    - Zero alterations to font.size, font.name, or line spacing.
    - Zero mentions of Streamlit or port 8501.
    - Exactly 6 slides (pruning Slide 7).
    - Embedding workflow_architecture.png on Slide 3 (Index 2).
    - Table on Slide 4 (Index 3).
    - Team Name 'CodeVanta' across all slides.
    """
    if not os.path.exists(input_pptx):
        raise FileNotFoundError(f"Template '{input_pptx}' not found!")

    prs = pptx.Presentation(input_pptx)
    print(f"Loaded '{input_pptx}' ({len(prs.slides)} slides).")

    # 1. Prune Slide 7 (Instructions / Excess slides) so slide count is strictly 6
    for i in range(len(prs.slides) - 1, -1, -1):
        sl = prs.slides[i]
        is_instruction = False
        for sp in sl.shapes:
            if sp.has_text_frame and ("IMPORTANT INSTRUCTIONS" in sp.text_frame.text or "Kindly keep the maximum slides" in sp.text_frame.text):
                is_instruction = True
                break
        if is_instruction or i >= 6:
            print(f"Pruning slide at index {i} (Instruction / Excess slide)...")
            rId = prs.slides._sldIdLst[i].rId
            prs.part.drop_rel(rId)
            del prs.slides._sldIdLst[i]

    assert len(prs.slides) == 6, f"Expected 6 slides, got {len(prs.slides)}"
    print("Slide pruning complete: Exactly 6 slides verified.")

    # Palette
    ROYAL_BLUE = RGBColor(29, 91, 165)  # #1d5ba5
    WHITE = RGBColor(255, 255, 255)
    ROW_BG = RGBColor(248, 250, 252)

    # =========================================================================
    # SLIDE 1 (INDEX 0): TITLE PAGE
    # =========================================================================
    s1 = prs.slides[0]
    for shape in s1.shapes:
        if shape.name == "Title 7" or (shape.has_text_frame and "SMART INDIA" in shape.text_frame.text):
            shape.text_frame.paragraphs[0].text = "SMART INDIA HACKATHON 2026"

        elif shape.name == "Subtitle 3" or (shape.has_text_frame and "TITLE PAGE" in shape.text_frame.text):
            for p in shape.text_frame.paragraphs:
                p.text = ""

        elif shape.has_text_frame and ("Problem Statement ID" in shape.text_frame.text or shape.name == "TextBox 9"):
            tf = shape.text_frame
            fields = [
                "Problem Statement ID: SIH26101",
                "Problem Statement Title: Karmayogi Statistical Intelligence (KSI) — AI-Powered Competency Modeling for MoSPI",
                "Theme: Smart Governance & Capacity Building (Mission Karmayogi Bharat)",
                "PS Category: Software (Air-Gapped Sovereign AI)",
                "Team Name: CodeVanta"
            ]
            # Assign directly to existing paragraphs to preserve layout/font inheritance
            for idx, text_val in enumerate(fields):
                if idx < len(tf.paragraphs):
                    tf.paragraphs[idx].text = text_val
                else:
                    p = tf.add_paragraph()
                    p.text = text_val
            # Clear any extra trailing paragraphs
            for idx in range(len(fields), len(tf.paragraphs)):
                tf.paragraphs[idx].text = ""

    # =========================================================================
    # SLIDE 2 (INDEX 1): PROPOSED SOLUTION
    # =========================================================================
    s2 = prs.slides[1]
    for shape in s2.shapes:
        if shape.has_text_frame:
            if shape.name == "Title 1" or "IDEA TITLE" in shape.text_frame.text:
                shape.text_frame.paragraphs[0].text = "PROPOSED SOLUTION: KARMAYOGI STATISTICAL INTELLIGENCE (KSI)"

            elif shape.name == "TextBox 8" or "Proposed Solution" in shape.text_frame.text:
                tf = shape.text_frame
                content_s2 = [
                    "Objective: Transitioning MoSPI from 'Rule-based' to 'Role-based' Competency Governance",
                    "• Deterministic Vector Gap Analysis: Models officer proficiency across 4 canonical domains against statutory benchmarks using exact L2 Euclidean distance.",
                    "• Deficit-Prioritized iGOT Pathways: Dynamic recommendation engine sorting accredited NSSTA/MeitY modules strictly by highest operational deficit.",
                    "• Air-Gapped Assessment Synthesizer: Native PyMuPDF parser + local LLM pipeline synthesizing Bloom-tiered MCQs directly from Gazette notifications.",
                    "• Sovereign Cadre Governance: 100% on-premise execution guaranteeing DPDP Act 2023 compliance, paired with Directorate capability heatmaps (NAD, FOD, PSD)."
                ]
                for idx, text_val in enumerate(content_s2):
                    if idx < len(tf.paragraphs):
                        tf.paragraphs[idx].text = text_val
                    else:
                        p = tf.add_paragraph()
                        p.text = text_val
                for idx in range(len(content_s2), len(tf.paragraphs)):
                    tf.paragraphs[idx].text = ""

    # =========================================================================
    # SLIDE 3 (INDEX 2): TECHNICAL APPROACH & ARCHITECTURE (2-COLUMN)
    # =========================================================================
    s3 = prs.slides[2]
    for shape in s3.shapes:
        if shape.has_text_frame and (shape.name == "Title 1" or "TECHNICAL APPROACH" in shape.text_frame.text):
            shape.text_frame.paragraphs[0].text = "TECHNICAL APPROACH & SYSTEM ARCHITECTURE"

    for shape in s3.shapes:
        if shape.has_text_frame and (shape.name == "TextBox 8" or "Technologies to be used" in shape.text_frame.text):
            shape.left = Inches(0.8)
            shape.top = Inches(1.8)
            shape.width = Inches(5.2)
            shape.height = Inches(4.8)

            tf = shape.text_frame
            bullets_s3 = [
                "• Core Stack: React 18 + Vite (Official iGOT Bharat Theme) | FastAPI Async REST | SQLite3 disk persistence (ksi_master.db).",
                "• Sovereign AI Engine: Local Ollama (qwen2.5:7b-instruct) pinned with 12s timeout & max_retries=0 anti-freeze guard. Zero external cloud API calls.",
                "• Vector Algebra: Exact Euclidean Deficit L2 = sqrt( sum( max(0, Benchmark_i - Score_i)^2 ) ).",
                "• Operational State: Enterprise web portal verified and live on client port 5173 (React/Vite) backed by async REST endpoints on port 8000 (FastAPI)."
            ]
            for idx, text_val in enumerate(bullets_s3):
                if idx < len(tf.paragraphs):
                    tf.paragraphs[idx].text = text_val
                else:
                    p = tf.add_paragraph()
                    p.text = text_val
            for idx in range(len(bullets_s3), len(tf.paragraphs)):
                tf.paragraphs[idx].text = ""

    # Remove any existing body image on Slide 3
    for sp in list(s3.shapes):
        if sp.shape_type == 13 and sp.top > Inches(1.0):
            elem = sp._element
            elem.getparent().remove(elem)

    # Embed workflow_architecture.png on Slide 3 right side
    if os.path.exists(arch_img):
        s3.shapes.add_picture(arch_img, Inches(6.2), Inches(1.8), width=Inches(6.6))
        print(f"[Slide 3] Embedded architecture image at Left 6.2\", Top 1.8\", Width 6.6\".")

    # =========================================================================
    # SLIDE 4 (INDEX 3): FEASIBILITY, RISK ANALYSIS & MITIGATION (TABLE)
    # =========================================================================
    s4 = prs.slides[3]
    for shape in s4.shapes:
        if shape.has_text_frame and (shape.name == "Title 1" or "FEASIBILITY" in shape.text_frame.text):
            shape.text_frame.paragraphs[0].text = "FEASIBILITY, RISK ANALYSIS & MITIGATION"

    # Clean up: Remove any placeholder text box or existing table from Slide 4
    for sp in list(s4.shapes):
        if sp.shape_type == 13 and sp.top > Inches(1.0):
            elem = sp._element
            elem.getparent().remove(elem)
        elif sp.has_table:
            elem = sp._element
            elem.getparent().remove(elem)
        elif sp.has_text_frame and (sp.name == "TextBox 8" or "Analysis of the feasibility" in sp.text_frame.text):
            elem = sp._element
            elem.getparent().remove(elem)

    # Add 4-row x 3-column Table on Slide 4
    tbl_shape = s4.shapes.add_table(
        rows=4,
        cols=3,
        left=Inches(0.8),
        top=Inches(1.8),
        width=Inches(11.6),
        height=Inches(4.5)
    )
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(2.3)
    tbl.columns[1].width = Inches(4.3)
    tbl.columns[2].width = Inches(5.0)

    headers = ["Dimension", "Potential Risk / Bottleneck", "Engineered Mitigation Strategy"]
    table_data = [
        ["Deployment Feasibility",
         "Hardware resource constraints on edge field devices without discrete GPUs.",
         "Engineered for commodity CPUs using quantized 7B model + zero-latency precomputed MD5 hash demo caching."],
        ["Inference Stability",
         "Local LLM latency spikes causing HTTP timeouts and interface freezes.",
         "Calibrated 12.0s ceiling with max_retries=0 fail-fast protection and deterministic pedagogical MCQ fallback generator."],
        ["Data Sovereignty & Legal",
         "Transmission of MoSPI employee records to third-party cloud APIs violates DPDP 2023.",
         "Zero-leak air-gapped architecture. All vector calculations, scores, and dossier embeddings execute strictly on localhost."]
    ]

    for c_i, h_txt in enumerate(headers):
        cell = tbl.cell(0, c_i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ROYAL_BLUE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.text_frame.paragraphs[0].text = h_txt
        # Set text color white for visibility on blue background (no font size/family overrides)
        cell.text_frame.paragraphs[0].font.color.rgb = WHITE
        cell.text_frame.paragraphs[0].font.bold = True

    for r_i, r_data in enumerate(table_data, start=1):
        bg = ROW_BG if (r_i % 2 == 1) else WHITE
        for c_i, val in enumerate(r_data):
            cell = tbl.cell(r_i, c_i)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.paragraphs[0].text = val
            if c_i == 0:
                cell.text_frame.paragraphs[0].font.bold = True
                cell.text_frame.paragraphs[0].font.color.rgb = ROYAL_BLUE

    print(f"[Slide 4] Created 4x3 Feasibility table (NO image on Slide 4).")

    # =========================================================================
    # SLIDE 5 (INDEX 4): INSTITUTIONAL IMPACT & NATIONAL BENEFITS
    # =========================================================================
    s5 = prs.slides[4]
    # Remove any tables or body images from Slide 5
    for sp in list(s5.shapes):
        if sp.has_table:
            elem = sp._element
            elem.getparent().remove(elem)
        elif sp.shape_type == 13 and sp.top > Inches(1.0):
            elem = sp._element
            elem.getparent().remove(elem)

    for shape in s5.shapes:
        if shape.has_text_frame:
            if shape.name == "Title 1" or "IMPACT" in shape.text_frame.text:
                shape.text_frame.paragraphs[0].text = "INSTITUTIONAL IMPACT & NATIONAL BENEFITS"

            elif shape.name == "TextBox 8" or "Potential impact" in shape.text_frame.text:
                tf = shape.text_frame
                bullets_s5 = [
                    "• Direct Beneficiaries: Empowers SSS and ISS officers across FOD, NAD, and PSD divisions.",
                    "• Objective Career Pathways: Replaces subjective annual performance appraisals with real-time, empirical competency telemetry and mathematical gap metrics.",
                    "• Rapid Curriculum Dissemination: Reduces lag between statutory revisions (e.g., SNA 2008 base revisions, GCF updates) and field readiness from months to seconds via automated circular parsing.",
                    "• Strategic Cadre Allocation: Provides Cadre Controlling Authorities with Directorate heatmaps to identify systemic deficits and dispatch targeted capacity-building directives.",
                    "• Viksit Bharat Alignment: Directly institutionalizes Mission Karmayogi's mandate for data-driven, role-competent public administration."
                ]
                for idx, text_val in enumerate(bullets_s5):
                    if idx < len(tf.paragraphs):
                        tf.paragraphs[idx].text = text_val
                    else:
                        p = tf.add_paragraph()
                        p.text = text_val
                for idx in range(len(bullets_s5), len(tf.paragraphs)):
                    tf.paragraphs[idx].text = ""

    print(f"[Slide 5] Populated 5 Impact bullets (NO table, NO image).")

    # =========================================================================
    # SLIDE 6 (INDEX 5): STATUTORY REFERENCES & RESEARCH WORK
    # =========================================================================
    s6 = prs.slides[5]
    for sp in list(s6.shapes):
        if sp.has_table:
            elem = sp._element
            elem.getparent().remove(elem)
        elif sp.shape_type == 13 and sp.top > Inches(1.0):
            elem = sp._element
            elem.getparent().remove(elem)

    for shape in s6.shapes:
        if shape.has_text_frame:
            if shape.name == "Title 1" or "RESEARCH" in shape.text_frame.text:
                shape.text_frame.paragraphs[0].text = "STATUTORY REFERENCES & RESEARCH WORK"

            elif shape.name == "TextBox 8" or "Details / Links" in shape.text_frame.text:
                tf = shape.text_frame
                bullets_s6 = [
                    "• Mission Karmayogi Bharat (DoPT / CBC): Framework of Roles, Activities, and Competencies (FRAC) Dictionary and NPCSCB guidelines.",
                    "• MoSPI / NSSTA Curriculum: Training Policy and Assessment Compendium (TPAC) and induction course syllabi for Junior Statistical Officers (JSO).",
                    "• United Nations Statistics Division: System of National Accounts (SNA 2008) guidelines on Gross Value Added (GVA), deflators, and input-output tables.",
                    "• Digital Personal Data Protection Act (DPDP 2023): Statutory provisions governing sovereign employee data localization, confidentiality, and consent boundaries.",
                    "• Pedagogical Grounding: Bloom's Revised Taxonomy (Anderson & Krathwohl, 2001) for automated multi-tier psychometric question generation."
                ]
                for idx, text_val in enumerate(bullets_s6):
                    if idx < len(tf.paragraphs):
                        tf.paragraphs[idx].text = text_val
                    else:
                        p = tf.add_paragraph()
                        p.text = text_val
                for idx in range(len(bullets_s6), len(tf.paragraphs)):
                    tf.paragraphs[idx].text = ""

    print(f"[Slide 6] Populated 5 Statutory Reference citations.")

    # =========================================================================
    # TEAM BRANDING: "CodeVanta" ACROSS ALL SLIDES
    # =========================================================================
    for idx, sl in enumerate(prs.slides):
        for sp in sl.shapes:
            if sp.has_text_frame:
                txt = sp.text_frame.text
                if "Your Team Name" in txt or "Runtime terror" in txt or "CodeVanta" in txt:
                    for p in sp.text_frame.paragraphs:
                        if "Your Team Name" in p.text or "Runtime terror" in p.text:
                            p.text = p.text.replace("Your Team Name", "CodeVanta").replace("Runtime terror", "CodeVanta")
                        if "CodeVanta" in p.text:
                            p.font.bold = True
                            p.font.color.rgb = ROYAL_BLUE

    # Save to primary and mirror
    prs.save(output_pptx)
    print(f"[SUCCESS] Saved primary output to '{output_pptx}'")

    mirror_pptx = "SIH2026_KSI_Submission_3.pptx"
    prs.save(mirror_pptx)
    print(f"[SUCCESS] Saved mirror output to '{mirror_pptx}'")

    mirror_2 = "SIH2026_KSI_Submission_2.pptx"
    prs.save(mirror_2)
    print(f"[SUCCESS] Saved mirror output to '{mirror_2}'")

if __name__ == '__main__':
    build_presentation()
