import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def fix_everything():
    # ALWAYS load original clean template to prevent compounding offset errors
    prs = Presentation("SIH2026-IDEA-Presentation-Format.pptx")
    
    # 1. Slide 1 (Title)
    s1 = prs.slides[0]
    for sp in s1.shapes:
        if sp.has_text_frame and "Problem Statement ID" in sp.text_frame.text:
            tf = sp.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "SMART INDIA HACKATHON 2026"
            p.font.bold = True
            p.font.size = Pt(22)
            p.font.color.rgb = RGBColor(29, 91, 165)
            
            items = [
                ("Problem Statement ID: ", "SIH26101"),
                ("Problem Statement Title: ", "Karmayogi Statistical Intelligence (KSI) — AI-Powered Competency Modeling for MoSPI"),
                ("Theme: ", "Smart Governance & Capacity Building (Mission Karmayogi Bharat)"),
                ("PS Category: ", "Software (Air-Gapped Sovereign AI)"),
                ("Team Name: ", "CodeVanta")
            ]
            for lbl, val in items:
                pi = tf.add_paragraph()
                r1 = pi.add_run()
                r1.text = lbl
                r1.font.bold = True
                r1.font.size = Pt(13)
                r2 = pi.add_run()
                r2.text = val
                r2.font.size = Pt(13)

    # 2. Slide 2 (Proposed Solution)
    s2 = prs.slides[1]
    for sp in s2.shapes:
        if sp.has_text_frame:
            if "IDEA TITLE" in sp.text_frame.text:
                sp.text_frame.paragraphs[0].text = "PROPOSED SOLUTION: KARMAYOGI STATISTICAL INTELLIGENCE (KSI)"
            elif "Proposed Solution" in sp.text_frame.text:
                tf = sp.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                p.text = "Objective: Transitioning MoSPI from 'Rule-based' to 'Role-based' Competency Governance"
                p.font.bold = True
                p.font.color.rgb = RGBColor(245, 130, 32)
                
                bullets = [
                    ("Deterministic Vector Gap Analysis: ", "Models officer proficiency across 4 canonical domains against statutory benchmarks using exact L2 Euclidean distance."),
                    ("Deficit-Prioritized iGOT Pathways: ", "Dynamic recommendation engine sorting accredited NSSTA/MeitY modules strictly by highest operational deficit."),
                    ("Air-Gapped Assessment Synthesizer: ", "Native PyMuPDF parser + local LLM pipeline synthesizing Bloom-tiered MCQs directly from Gazette notifications."),
                    ("Sovereign Cadre Governance: ", "100% on-premise execution guaranteeing DPDP Act 2023 compliance, paired with Directorate capability heatmaps (NAD, FOD, PSD).")
                ]
                for b_title, b_desc in bullets:
                    p_b = tf.add_paragraph()
                    r1 = p_b.add_run()
                    r1.text = "• " + b_title
                    r1.font.bold = True
                    r2 = p_b.add_run()
                    r2.text = b_desc

    # 3. Slide 3 (Technical Approach + Image side-by-side)
    s3 = prs.slides[2]
    for sp in s3.shapes:
        if sp.has_text_frame:
            if "TECHNICAL APPROACH" in sp.text_frame.text:
                sp.text_frame.paragraphs[0].text = "TECHNICAL APPROACH & SYSTEM ARCHITECTURE"
            elif "Technologies to be used" in sp.text_frame.text:
                sp.left = Inches(0.8)
                sp.top = Inches(1.8)
                sp.width = Inches(5.2)
                sp.height = Inches(4.8)
                tf = sp.text_frame
                tf.clear()
                
                specs = [
                    ("Core Stack: ", "React 18 + Vite (Official iGOT Bharat Theme) | FastAPI Async REST | SQLite3 disk persistence (ksi_master.db)."),
                    ("Sovereign AI Engine: ", "Local Ollama (qwen2.5:7b-instruct) pinned with 12s timeout & max_retries=0 anti-freeze guard. Zero external cloud API calls."),
                    ("Vector Algebra: ", "Exact Euclidean Deficit L2 = sqrt( sum( max(0, Benchmark_i - Score_i)^2 ) )."),
                    ("Operational State: ", "Enterprise web portal verified and live on client port 5173 (React/Vite) backed by async REST endpoints on port 8000 (FastAPI).")
                ]
                for i, (st, sd) in enumerate(specs):
                    p_s = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    r1 = p_s.add_run()
                    r1.text = "• " + st
                    r1.font.bold = True
                    r2 = p_s.add_run()
                    r2.text = sd

    if os.path.exists("workflow_architecture.png"):
        s3.shapes.add_picture("workflow_architecture.png", Inches(6.2), Inches(1.8), width=Inches(6.5))

    # 4. Slide 4 (Feasibility Table ONLY)
    s4 = prs.slides[3]
    for sp in list(s4.shapes):
        if sp.has_text_frame:
            if "FEASIBILITY" in sp.text_frame.text:
                sp.text_frame.paragraphs[0].text = "FEASIBILITY, RISK ANALYSIS & MITIGATION"
            else:
                sp_elem = sp._element
                sp_elem.getparent().remove(sp_elem)

    tbl_shape = s4.shapes.add_table(4, 3, Inches(0.8), Inches(1.8), Inches(11.6), Inches(4.5))
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(2.8)
    tbl.columns[1].width = Inches(4.2)
    tbl.columns[2].width = Inches(4.6)

    headers = ["Dimension", "Potential Risk / Bottleneck", "Engineered Mitigation Strategy"]
    for idx, h in enumerate(headers):
        cell = tbl.cell(0, idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(29, 91, 165)
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)

    data = [
        ("Deployment Feasibility", "Hardware resource constraints on edge field devices without discrete GPUs.", "Engineered for commodity CPUs using quantized 7B model + zero-latency precomputed MD5 hash demo caching."),
        ("Inference Stability", "Local LLM latency spikes causing HTTP timeouts and interface freezes.", "Calibrated 12.0s ceiling with max_retries=0 fail-fast protection and deterministic pedagogical MCQ fallback generator."),
        ("Data Sovereignty & Legal", "Transmission of MoSPI employee records to third-party cloud APIs violates DPDP 2023.", "Zero-leak air-gapped architecture. All vector calculations, scores, and dossier embeddings execute strictly on localhost.")
    ]
    for r_idx, row in enumerate(data, start=1):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(248, 250, 252) if r_idx % 2 == 0 else RGBColor(255, 255, 255)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10)
            p.font.color.rgb = RGBColor(30, 41, 59)

    # 5. Slide 5 (Impact)
    s5 = prs.slides[4]
    for sp in s5.shapes:
        if sp.has_text_frame:
            if "IMPACT" in sp.text_frame.text:
                sp.text_frame.paragraphs[0].text = "INSTITUTIONAL IMPACT & NATIONAL BENEFITS"
            elif "Potential impact" in sp.text_frame.text:
                tf = sp.text_frame
                tf.clear()
                benefits = [
                    ("Direct Beneficiaries: ", "Empowers Subordinate Statistical Service (SSS) and Indian Statistical Service (ISS) officers across FOD, NAD, and PSD divisions."),
                    ("Objective Career Pathways: ", "Replaces subjective annual performance appraisals with real-time, empirical competency telemetry and mathematical gap metrics."),
                    ("Rapid Curriculum Dissemination: ", "Reduces lag between statutory revisions (e.g., SNA 2008 base revisions, GCF updates) and field readiness from months to seconds via automated circular parsing."),
                    ("Strategic Cadre Allocation: ", "Provides Cadre Controlling Authorities with Directorate heatmaps to identify systemic deficits and dispatch targeted capacity-building directives."),
                    ("Viksit Bharat Alignment: ", "Directly institutionalizes Mission Karmayogi's mandate for data-driven, role-competent public administration.")
                ]
                for i, (bt, bd) in enumerate(benefits):
                    p_b = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    r1 = p_b.add_run()
                    r1.text = "• " + bt
                    r1.font.bold = True
                    r2 = p_b.add_run()
                    r2.text = bd

    # 6. Slide 6 (References)
    s6 = prs.slides[5]
    for sp in s6.shapes:
        if sp.has_text_frame:
            if "RESEARCH" in sp.text_frame.text:
                sp.text_frame.paragraphs[0].text = "STATUTORY REFERENCES & RESEARCH WORK"
            elif "Details / Links" in sp.text_frame.text:
                tf = sp.text_frame
                tf.clear()
                refs = [
                    ("Mission Karmayogi Bharat (DoPT / CBC): ", "Framework of Roles, Activities, and Competencies (FRAC) Dictionary and NPCSCB guidelines."),
                    ("MoSPI / NSSTA Curriculum: ", "Training Policy and Assessment Compendium (TPAC) and induction course syllabi for Junior Statistical Officers (JSO)."),
                    ("United Nations Statistics Division: ", "System of National Accounts (SNA 2008) guidelines on Gross Value Added (GVA), deflators, and input-output tables."),
                    ("Digital Personal Data Protection Act (DPDP 2023): ", "Statutory provisions governing sovereign employee data localization, confidentiality, and consent boundaries."),
                    ("Pedagogical Grounding: Bloom's Revised Taxonomy (Anderson & Krathwohl, 2001) for automated multi-tier psychometric question generation.")
                ]
                for i, r in enumerate(refs):
                    p_r = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    r1 = p_r.add_run()
                    r1.text = "• " + (r[0] if isinstance(r, tuple) else r)
                    r1.font.bold = True
                    if isinstance(r, tuple):
                        r2 = p_r.add_run()
                        r2.text = r[1]

    # Prune Slide 7
    while len(prs.slides) > 6:
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]

    # Global footer sync
    for slide in prs.slides:
        for sp in slide.shapes:
            if sp.has_text_frame and ("Your Team Name" in sp.text_frame.text or "CodeVanta" in sp.text_frame.text):
                sp.text_frame.text = sp.text_frame.text.replace("Your Team Name", "CodeVanta")

    # Save to requested filename and parity filenames
    prs.save("SIH2026_KSI_Submission_FINAL.pptx")
    print("Clean 6-slide deck generated successfully as SIH2026_KSI_Submission_FINAL.pptx")
    
    prs.save("SIH2026_KSI_Submission.pptx")
    print("Synchronized clean 6-slide deck to SIH2026_KSI_Submission.pptx")

if __name__ == "__main__":
    fix_everything()
